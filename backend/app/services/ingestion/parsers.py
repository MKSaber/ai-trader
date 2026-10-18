"""CSV/XLSX market-data parsers with column-mapping preview."""

from __future__ import annotations

import csv
import io
from datetime import date, datetime, timezone
from pathlib import Path
from typing import Any

from app.ports.market_data import BarRow

CANONICAL = {
    "symbol": ("symbol", "ticker", "yahoo_symbol"),
    "session_date": ("session_date", "date", "timestamp"),
    "open": ("open", "open_raw"),
    "high": ("high", "high_raw"),
    "low": ("low", "low_raw"),
    "close": ("close", "close_raw"),
    "volume": ("volume",),
    "currency": ("currency",),
}


def preview_columns(data: bytes, filename: str) -> dict[str, Any]:
    """Return detected headers, suggested mapping, and a short preview."""
    rows, headers = _load_rows(data, filename)
    lower = {h.lower(): h for h in headers}
    mapping: dict[str, str | None] = {}
    for canon, aliases in CANONICAL.items():
        mapping[canon] = None
        for alias in aliases:
            if alias in lower:
                mapping[canon] = lower[alias]
                break
    missing = [k for k, v in mapping.items() if v is None and k != "currency" and k != "volume"]
    return {
        "headers": headers,
        "suggested_mapping": mapping,
        "missing_required": missing,
        "preview_rows": rows[:5],
        "row_count_estimate": len(rows),
    }


def parse_market_rows(
    data: bytes,
    filename: str,
    column_mapping: dict[str, str] | None = None,
) -> list[BarRow]:
    """Parse market OHLCV from CSV or XLSX into BarRow list."""
    rows, headers = _load_rows(data, filename)
    if not headers:
        raise ValueError("File has no header row")
    lower = {h.lower(): h for h in headers}
    mapping = column_mapping or {}
    resolved: dict[str, str] = {}
    for canon, aliases in CANONICAL.items():
        if mapping.get(canon):
            resolved[canon] = mapping[canon]
            continue
        for alias in aliases:
            if alias in lower:
                resolved[canon] = lower[alias]
                break
    for required in ("symbol", "session_date", "open", "high", "low", "close"):
        if required not in resolved:
            raise ValueError(
                f"Missing required column for {required}. "
                f"Headers={headers}"
            )

    bars: list[BarRow] = []
    for row in rows:
        raw_date = str(row[resolved["session_date"]])
        if "T" in raw_date:
            ts = datetime.fromisoformat(raw_date.replace("Z", "+00:00"))
            session = ts.date()
        else:
            # Excel may give datetime objects already stringified
            session = date.fromisoformat(raw_date[:10])
            ts = datetime(
                session.year, session.month, session.day, tzinfo=timezone.utc
            )
        sym = str(row[resolved["symbol"]]).strip()
        if not sym.endswith(".L") and "." not in sym:
            sym = f"{sym}.L"
        currency = "GBX"
        if "currency" in resolved and row.get(resolved["currency"]):
            currency = str(row[resolved["currency"]])
        vol = 0.0
        if "volume" in resolved and row.get(resolved["volume"]) not in (None, ""):
            vol = float(row[resolved["volume"]])
        bars.append(
            BarRow(
                symbol=sym,
                bar_ts=ts,
                session_date=session,
                open=float(row[resolved["open"]]),
                high=float(row[resolved["high"]]),
                low=float(row[resolved["low"]]),
                close=float(row[resolved["close"]]),
                volume=vol,
                currency_hint=currency,
            )
        )
    return bars


def _load_rows(data: bytes, filename: str) -> tuple[list[dict[str, Any]], list[str]]:
    """Load tabular rows from CSV or XLSX bytes."""
    suffix = Path(filename).suffix.lower()
    if suffix in {".xlsx", ".xls"}:
        return _load_xlsx(data)
    text = data.decode("utf-8", errors="replace")
    reader = csv.DictReader(io.StringIO(text))
    if not reader.fieldnames:
        return [], []
    rows = [dict(r) for r in reader]
    return rows, list(reader.fieldnames)


def _load_xlsx(data: bytes) -> tuple[list[dict[str, Any]], list[str]]:
    """Load first sheet of an XLSX workbook."""
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    ws = wb.active
    it = ws.iter_rows(values_only=True)
    try:
        header_row = next(it)
    except StopIteration:
        return [], []
    headers = [str(h).strip() if h is not None else f"col_{i}" for i, h in enumerate(header_row)]
    rows: list[dict[str, Any]] = []
    for values in it:
        if all(v is None for v in values):
            continue
        row = {}
        for h, v in zip(headers, values):
            if isinstance(v, datetime):
                row[h] = v.isoformat()
            elif isinstance(v, date):
                row[h] = v.isoformat()
            else:
                row[h] = v
        rows.append(row)
    return rows, headers
